"""FINAL PCN874 input-VAT file (ready for reporting).

Combines BOTH datasets (approved expenses + pending drafts), keeps only invoices
with deductible Israeli VAT (vat>0, ILS), from 2025-06-01 onward, DEDUPLICATES
across/within datasets, computes suggested deductible VAT (full / 2/3 vehicle /
0 כיבוד), and lays it out by bi-monthly period — ready for the client's review.
"""
import re
from collections import defaultdict
from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

GREEN = PatternFill("solid", fgColor="C6EFCE"); YELLOW = PatternFill("solid", fgColor="FFEB9C")
RED = PatternFill("solid", fgColor="FFC7CE"); HDR = PatternFill("solid", fgColor="305496")
SUB = PatternFill("solid", fgColor="DDEBF7"); THIN = Border(*[Side(style="thin", color="D9D9D9")]*4)
CF = {"high": GREEN, "medium": YELLOW, "low": RED}
HEBMON = {"01":"ינו","02":"פבר","03":"מרץ","04":"אפר","05":"מאי","06":"יוני","07":"יולי","08":"אוג","09":"ספט","10":"אוק","11":"נוב","12":"דצמ"}


def fnum(v):
    try: return round(float(v), 2)
    except Exception: return None


def bimonthly(d):
    m = re.search(r"(\d{4})-(\d{2})", str(d or ""))
    if not m: return ""
    y, mo = m.group(1), int(m.group(2)); s = mo if mo % 2 == 1 else mo-1
    return f"{s:02d}-{s+1:02d}/{y}"


def conf_norm(v):
    v = str(v or "")
    if "גבוה" in v or "🟩" in v: return "high"
    if "בינ" in v or "🟨" in v: return "medium"
    if "נמוכ" in v or "🟥" in v or "לבדיקה" in v: return "low"
    return ""


def load_task1():
    ws = load_workbook("catalog.xlsx")["קטלוג 441"]; h = [c.value for c in ws[1]]; c = lambda n: h.index(n)
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[c("#index")] in (None, ""): continue
        out.append({"src":"טיוטה","sup":r[c("שם ספק")],"tax":r[c("מספר עוסק / ח.פ.")],"dtype":r[c("סוג מסמך")],
                    "num":r[c("מספר מסמך")],"date":r[c("תאריך המסמך")],"amount":fnum(r[c("סכום כולל מע\"מ")]),
                    "vat":fnum(r[c("סך הכל מע\"מ")]),"cur":r[c("מטבע")] or "ILS","et":r[c("סוג הוצאה (מוצע)")],
                    "pcn":r[c("סיווג PCN874")],"conf":conf_norm(r[c("רמת ודאות כללית")]),"note":r[c("הערות")],
                    "id":r[c("draft_id")],"file":r[c("שם הקובץ")]})
    return out


def load_task2():
    ws = load_workbook("catalog_expenses_FINAL.xlsx")["פירוט לפי חודש"]; h = [c.value for c in ws[1]]; c = lambda n: h.index(n)
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[c("#")] in (None,"") or str(r[c("expense_id")] or "")=="": continue
        out.append({"src":"מאושר","sup":r[c("שם ספק")],"tax":r[c("ע.מ/ח.פ")],"dtype":r[c("סוג מסמך")],
                    "num":r[c("מס' מסמך")],"date":r[c("תאריך")],"amount":fnum(r[c("כולל מע\"מ")]),
                    "vat":fnum(r[c("מע\"מ")]),"cur":r[c("מטבע")] or "ILS","et":r[c("סוג הוצאה")],
                    "pcn":r[c("סיווג PCN874")],"conf":conf_norm(r[c("ודאות")]),"note":r[c("שאלה/ספק לליבון")],
                    "id":r[c("expense_id")],"file":r[c("קובץ")]})
    return out


def deductible(et, vat):
    et = et or ""
    if any(k in et for k in ["דלק","אחזקת רכב","רכב"]): return round(vat*2/3,2), "רכב 2/3"
    if "כיבוד" in et: return 0, "כיבוד — לא מוכר (לאישור)"
    return vat, ""


def sig(d):
    t = re.sub(r"\D","",str(d["tax"] or "")); n = re.sub(r"\W","",str(d["num"] or "")).lower()
    a = str(d["amount"])
    if n and len(n) >= 4: return f"{t}|{n}|{a}"
    return f"{(d['sup'] or '')[:12]}|{a}|{d['date']}"


COLS = ["מקור","תאריך","תקופה","שם ספק","ע.מ/ח.פ","סוג מסמך","מס' מסמך","סכום כולל",
        "מע\"מ בחשבונית","מע\"מ לקיזוז (מוצע)","סיווג PCN874","סוג הוצאה","ודאות","סטטוס","הערה","קובץ"]


def main():
    rows = load_task1() + load_task2()
    # keep VAT>0, ILS, date>=2025-06
    rows = [d for d in rows if (d["vat"] or 0) > 0 and (d["cur"] or "ILS") == "ILS"
            and str(d["date"] or "") >= "2025-06"]
    # dedup
    seen = {}
    for d in rows:
        s = sig(d); d["dup"] = False
        if s in seen:
            d["dup"] = True; d["dupof"] = seen[s]
        else:
            seen[s] = (d["id"], d["src"])
    rows.sort(key=lambda d: (bimonthly(d["date"]), d["date"], d["sup"] or ""))

    wb = Workbook()
    sm = wb.active; sm.title = "סיכום לתקופות"; sm.sheet_view.rightToLeft = True
    sm.append(["תקופה דו-חודשית","חשבוניות (ללא כפולות)","מע\"מ בחשבוניות","מע\"מ לקיזוז (מוצע)"])
    for cc in sm[1]: cc.fill=HDR; cc.font=Font(bold=True,color="FFFFFF"); cc.border=THIN
    per = defaultdict(lambda:[0,0.0,0.0])
    for d in rows:
        if d["dup"]: continue
        ded,_ = deductible(d["et"], d["vat"]); p = per[bimonthly(d["date"])]
        p[0]+=1; p[1]+=d["vat"]; p[2]+=ded
    for k in sorted(per):
        v=per[k]; sm.append([k,v[0],round(v[1],2),round(v[2],2)])
    sm.append([])
    sm.append(["סה\"כ",sum(v[0] for v in per.values()),round(sum(v[1] for v in per.values()),2),round(sum(v[2] for v in per.values()),2)])
    for cc in sm[sm.max_row]: cc.font=Font(bold=True); cc.fill=SUB
    for i,w in enumerate([20,20,18,22],1): sm.column_dimensions[get_column_letter(i)].width=w

    ws = wb.create_sheet("חשבוניות לקיזוז"); ws.sheet_view.rightToLeft=True
    for cc,name in enumerate(COLS,1):
        cell=ws.cell(1,cc,name); cell.fill=HDR; cell.font=Font(bold=True,color="FFFFFF",size=10)
        cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); cell.border=THIN
    ws.freeze_panes="A2"; r=2; cur=None
    for d in rows:
        p=bimonthly(d["date"])
        if p!=cur:
            cur=p; ws.cell(r,1,f"— תקופה {p} —").font=Font(bold=True)
            for cc in range(1,len(COLS)+1): ws.cell(r,cc).fill=SUB
            r+=1
        ded,dnote=deductible(d["et"],d["vat"])
        status = "⚠️ כפילות — לא לקזז" if d["dup"] else ""
        vals={"מקור":d["src"],"תאריך":d["date"],"תקופה":p,"שם ספק":d["sup"],"ע.מ/ח.פ":d["tax"],
              "סוג מסמך":d["dtype"],"מס' מסמך":d["num"],"סכום כולל":d["amount"],"מע\"מ בחשבונית":d["vat"],
              "מע\"מ לקיזוז (מוצע)":(0 if d["dup"] else ded),"סיווג PCN874":d["pcn"],"סוג הוצאה":d["et"],
              "ודאות":{"high":"🟩","medium":"🟨","low":"🟥"}.get(d["conf"],""),"סטטוס":status,
              "הערה":((d["note"] or "")+((" | "+dnote) if dnote else "")).strip(" |"),"קובץ":d["file"]}
        for cc,name in enumerate(COLS,1):
            cell=ws.cell(r,cc,vals[name]); cell.border=THIN
            cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
            if name=="ודאות" and d["conf"] in CF: cell.fill=CF[d["conf"]]
            if d["dup"]: cell.fill=RED
        r+=1
    for cc,w in enumerate([12,11,10,24,13,16,15,11,12,14,18,18,7,18,30,28],1):
        ws.column_dimensions[get_column_letter(cc)].width=w
    wb.save("PCN_FINAL.xlsx")
    nd=sum(1 for d in rows if not d["dup"]); dd=sum(1 for d in rows if d["dup"])
    print(f"PCN_FINAL: {len(rows)} VAT rows | unique={nd} dups={dd} | invoice VAT={round(sum(d['vat'] for d in rows if not d['dup']),2)} | deductible={round(sum(deductible(d['et'],d['vat'])[0] for d in rows if not d['dup']),2)}")


if __name__ == "__main__":
    main()
