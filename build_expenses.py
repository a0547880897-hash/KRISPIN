"""Task 2 builder: catalog_expenses_FINAL.xlsx from MY independent scan.

Reads expenses_scan.json (my per-document extraction) + index_expenses.csv.
Organised by month. Independent scan (no Morning comparison in the sheet).
Cells coloured by my confidence. Doubts flagged. Per-row hyperlink to source.
"""
import os, csv, json
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SCAN = "expenses_scan.json"; INDEX = "index_expenses.csv"
DOCS = "documents_expenses"; OUT = "catalog_expenses_FINAL.xlsx"
GREEN = PatternFill("solid", fgColor="C6EFCE"); YELLOW = PatternFill("solid", fgColor="FFEB9C")
RED = PatternFill("solid", fgColor="FFC7CE"); HDR = PatternFill("solid", fgColor="305496")
SUB = PatternFill("solid", fgColor="DDEBF7"); THIN = Border(*[Side(style="thin", color="D9D9D9")] * 4)
CF = {"high": GREEN, "medium": YELLOW, "low": RED}
HEBMON = {"01": "ינואר", "02": "פברואר", "03": "מרץ", "04": "אפריל", "05": "מאי", "06": "יוני",
          "07": "יולי", "08": "אוגוסט", "09": "ספטמבר", "10": "אוקטובר", "11": "נובמבר", "12": "דצמבר"}
COLS = ["#", "תאריך", "שם ספק", "ע.מ/ח.פ", "סוג מסמך", "מס' מסמך", "לפני מע\"מ", "מע\"מ",
        "כולל מע\"מ", "מטבע", "שולם", "סוג הוצאה", "סיווג PCN874", "ודאות", "שאלה/ספק לליבון", "קובץ", "expense_id"]


def fnum(v):
    try: return float(v)
    except Exception: return 0


def main():
    scan = json.load(open(SCAN, encoding="utf-8")) if os.path.exists(SCAN) else {}
    scan = {k: v for k, v in scan.items() if not k.startswith("_")}
    idx = {}
    if os.path.exists(INDEX):
        for r in csv.DictReader(open(INDEX, encoding="utf-8-sig")):
            idx[r["expense_id"]] = r
    files = {f.split("_")[1]: f for f in os.listdir(DOCS)} if os.path.isdir(DOCS) else {}
    rows = []
    for eid, meta in idx.items():
        s = scan.get(eid, {})
        date = s.get("date") or (meta.get("reportingDate") or meta.get("date") or "")[:10]
        rows.append(((date or "")[:7], eid, meta, s, date))
    rows.sort(key=lambda x: (x[0], x[1]))

    wb = Workbook()
    sm = wb.active; sm.title = "סיכום חודשי"; sm.sheet_view.rightToLeft = True
    sm.append(["חודש", "מס' מסמכים", "סה\"כ כולל מע\"מ", "סה\"כ מע\"מ תשומות", "נסרקו"])
    for c in sm[1]:
        c.fill = HDR; c.font = Font(bold=True, color="FFFFFF"); c.border = THIN
    bm = defaultdict(lambda: [0, 0.0, 0.0, 0])
    for ym, eid, meta, s, date in rows:
        b = bm[ym]; b[0] += 1; b[1] += fnum(s.get("amount")); b[2] += fnum(s.get("vat")); b[3] += 1 if s else 0
    for ym in sorted(bm):
        b = bm[ym]; lbl = f"{HEBMON.get(ym[5:7], ym[5:7])} {ym[:4]}" if len(ym) >= 7 else "ללא תאריך"
        sm.append([lbl, b[0], round(b[1], 2), round(b[2], 2), f"{b[3]}/{b[0]}"])
    sm.append([])
    sm.append(["סה\"כ", sum(b[0] for b in bm.values()), round(sum(b[1] for b in bm.values()), 2),
               round(sum(b[2] for b in bm.values()), 2),
               f"{sum(b[3] for b in bm.values())}/{sum(b[0] for b in bm.values())}"])
    for c in sm[sm.max_row]:
        c.font = Font(bold=True); c.fill = SUB
    for i, w in enumerate([16, 13, 18, 18, 10], 1):
        sm.column_dimensions[get_column_letter(i)].width = w

    ws = wb.create_sheet("פירוט לפי חודש"); ws.sheet_view.rightToLeft = True
    for c, name in enumerate(COLS, 1):
        cell = ws.cell(1, c, name); cell.fill = HDR; cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = THIN
    ws.freeze_panes = "A2"
    r = 2; cur = None; st = {"green": 0, "yellow": 0, "red": 0, "scanned": 0}
    for ym, eid, meta, s, date in rows:
        if ym != cur:
            cur = ym; lbl = f"— {HEBMON.get(ym[5:7], ym[5:7])} {ym[:4]} —" if len(ym) >= 7 else "— ללא תאריך —"
            ws.cell(r, 1, lbl).font = Font(bold=True)
            for c in range(1, len(COLS) + 1): ws.cell(r, c).fill = SUB
            r += 1
        if s: st["scanned"] += 1
        fname = files.get(eid, ""); conf = s.get("confidence", "")
        vals = {"#": meta.get("index", ""), "תאריך": date, "שם ספק": s.get("supplier", ""),
                "ע.מ/ח.פ": s.get("taxId", ""), "סוג מסמך": s.get("doc_type", ""), "מס' מסמך": s.get("number", ""),
                "לפני מע\"מ": s.get("net", ""), "מע\"מ": s.get("vat", ""), "כולל מע\"מ": s.get("amount", ""),
                "מטבע": s.get("currency", ""), "שולם": s.get("payment", ""), "סוג הוצאה": s.get("expense_type", ""),
                "סיווג PCN874": s.get("pcn", ""),
                "ודאות": {"high": "🟩 גבוהה", "medium": "🟨 בינונית", "low": "🟥 לבדיקה"}.get(conf, "טרם נסרק"),
                "שאלה/ספק לליבון": s.get("doubt", ""), "קובץ": fname, "expense_id": eid}
        for c, name in enumerate(COLS, 1):
            cell = ws.cell(r, c, vals[name]); cell.border = THIN
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if name == "ודאות" and conf in CF: cell.fill = CF[conf]
            if name == "קובץ" and fname:
                cell.hyperlink = DOCS + "/" + fname; cell.font = Font(color="0563C1", underline="single")
        st["green" if conf == "high" else "yellow" if conf == "medium" else "red" if conf == "low" else "scanned"] += 0
        if conf == "high": st["green"] += 1
        elif conf == "medium": st["yellow"] += 1
        elif conf == "low": st["red"] += 1
        r += 1
    for c, w in enumerate([6, 11, 24, 13, 20, 14, 11, 10, 11, 7, 16, 22, 22, 12, 34, 30, 24], 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    wb.save(OUT)
    print(f"Wrote {OUT}: {len(rows)} rows, scanned={st['scanned']} (G{st['green']} Y{st['yellow']} R{st['red']})")


if __name__ == "__main__":
    main()
