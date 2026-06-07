"""Build the PCN874 INPUT-VAT list: only expenses WITH deductible VAT.

Filters my scan (expenses_scan.json) to rows with vat>0, computes a suggested
deductible VAT (full for regular inputs, 2/3 for vehicle fuel/maintenance,
flagged for כיבוד), groups by bi-monthly period, and flags duplicates / items
needing the client's decision. Excludes no-VAT docs (Eilat, donations, arnona,
foreign services, tickets). Nothing touches Morning.
"""
import os, json, csv
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SCAN = "expenses_scan.json"; INDEX = "index_expenses.csv"; DOCS = "documents_expenses"
OUT = "PCN_input_VAT.xlsx"
GREEN = PatternFill("solid", fgColor="C6EFCE"); YELLOW = PatternFill("solid", fgColor="FFEB9C")
RED = PatternFill("solid", fgColor="FFC7CE"); HDR = PatternFill("solid", fgColor="305496")
SUB = PatternFill("solid", fgColor="DDEBF7"); THIN = Border(*[Side(style="thin", color="D9D9D9")]*4)
CF = {"high": GREEN, "medium": YELLOW, "low": RED}
HEBMON = {"01":"ינו","02":"פבר","03":"מרץ","04":"אפר","05":"מאי","06":"יוני","07":"יולי","08":"אוג","09":"ספט","10":"אוק","11":"נוב","12":"דצמ"}


def bimonthly(d):
    if not d or len(d) < 7: return "לא ידוע"
    y, m = d[:4], int(d[5:7]); s = m if m % 2 == 1 else m-1
    return f"{s:02d}-{s+1:02d}/{y}"


def fnum(v):
    try: return round(float(v), 2)
    except Exception: return None


def deductible(et, vat):
    """suggested deductible VAT + note."""
    et = et or ""
    if any(k in et for k in ["דלק", "רכב", "אחזקת רכב"]):
        return round(vat * 2/3, 2), "רכב — 2/3"
    if "כיבוד" in et:
        return 0, "כיבוד — בד\"כ לא מוכר (לאישורך)"
    if "מתנ" in et:
        return vat, "מתנה — בכפוף לתקרה"
    return vat, ""


COLS = ["#", "תאריך", "תקופה", "שם ספק", "ע.מ/ח.פ", "סוג מסמך", "מס' מסמך",
        "סכום כולל מע\"מ", "מע\"מ בחשבונית", "מע\"מ לקיזוז (מוצע)", "סיווג PCN874",
        "סוג הוצאה", "ודאות", "הערה/ספק", "קובץ"]


def main():
    scan = {k: v for k, v in json.load(open(SCAN, encoding="utf-8")).items() if not k.startswith("_")}
    files = {f.split("_")[1]: f for f in os.listdir(DOCS)} if os.path.isdir(DOCS) else {}
    idx = {r["expense_id"]: r["index"] for r in csv.DictReader(open(INDEX, encoding="utf-8-sig"))}

    rows = []
    for eid, s in scan.items():
        vat = fnum(s.get("vat"))
        if not vat or vat <= 0:
            continue  # only docs WITH input VAT
        cur = s.get("currency", "ILS")
        if cur != "ILS":
            continue  # foreign-currency services carry no Israeli input VAT
        ded, note = deductible(s.get("expense_type"), vat)
        date = s.get("date") or ""
        dup = "כפילות" in (s.get("doubt") or "")
        rows.append({"eid": eid, "idx": idx.get(eid, "?"), "date": date, "period": bimonthly(date),
                     "sup": s.get("supplier", ""), "tax": s.get("taxId", ""), "dtype": s.get("doc_type", ""),
                     "num": s.get("number", ""), "amount": fnum(s.get("amount")), "vat": vat,
                     "ded": ded, "pcn": s.get("pcn", ""), "et": s.get("expense_type", ""),
                     "conf": s.get("confidence", ""), "note": ((s.get("doubt") or "") + (" | " + note if note else "")).strip(" |"),
                     "dup": dup, "file": files.get(eid, "")})
    rows.sort(key=lambda r: (r["period"], r["date"], r["idx"]))

    wb = Workbook()
    # summary by period
    sm = wb.active; sm.title = "סיכום לתקופות"; sm.sheet_view.rightToLeft = True
    sm.append(["תקופה דו-חודשית", "מס' חשבוניות", "סה\"כ מע\"מ בחשבוניות", "סה\"כ מע\"מ לקיזוז (מוצע)"])
    for c in sm[1]:
        c.fill = HDR; c.font = Font(bold=True, color="FFFFFF"); c.border = THIN
    per = defaultdict(lambda: [0, 0.0, 0.0])
    for r in rows:
        p = per[r["period"]]; p[0]+=1; p[1]+=r["vat"]; p[2]+=r["ded"]
    for k in sorted(per):
        v = per[k]; sm.append([k, v[0], round(v[1],2), round(v[2],2)])
    sm.append([])
    sm.append(["סה\"כ", sum(v[0] for v in per.values()), round(sum(v[1] for v in per.values()),2), round(sum(v[2] for v in per.values()),2)])
    for c in sm[sm.max_row]: c.font = Font(bold=True); c.fill = SUB
    for i,w in enumerate([18,14,22,24],1): sm.column_dimensions[get_column_letter(i)].width=w

    ws = wb.create_sheet("חשבוניות לקיזוז"); ws.sheet_view.rightToLeft = True
    for c,name in enumerate(COLS,1):
        cell=ws.cell(1,c,name); cell.fill=HDR; cell.font=Font(bold=True,color="FFFFFF",size=10)
        cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); cell.border=THIN
    ws.freeze_panes="A2"; r=2; cur_p=None
    for d in rows:
        if d["period"]!=cur_p:
            cur_p=d["period"]; ws.cell(r,1,f"— תקופה {cur_p} —").font=Font(bold=True)
            for c in range(1,len(COLS)+1): ws.cell(r,c).fill=SUB
            r+=1
        vals={"#":d["idx"],"תאריך":d["date"],"תקופה":d["period"],"שם ספק":d["sup"],"ע.מ/ח.פ":d["tax"],
              "סוג מסמך":d["dtype"],"מס' מסמך":d["num"],"סכום כולל מע\"מ":d["amount"],"מע\"מ בחשבונית":d["vat"],
              "מע\"מ לקיזוז (מוצע)":d["ded"],"סיווג PCN874":d["pcn"],"סוג הוצאה":d["et"],
              "ודאות":{"high":"🟩","medium":"🟨","low":"🟥"}.get(d["conf"],""),"הערה/ספק":d["note"],"קובץ":d["file"]}
        for c,name in enumerate(COLS,1):
            cell=ws.cell(r,c,vals[name]); cell.border=THIN
            cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
            if name=="ודאות" and d["conf"] in CF: cell.fill=CF[d["conf"]]
            if d["dup"]: cell.fill=RED
            if name=="קובץ" and d["file"]:
                cell.hyperlink=DOCS+"/"+d["file"]; cell.font=Font(color="0563C1",underline="single")
        r+=1
    for c,w in enumerate([6,11,10,24,13,18,14,12,12,14,20,20,7,34,30],1):
        ws.column_dimensions[get_column_letter(c)].width=w
    wb.save(OUT)
    print(f"Wrote {OUT}: {len(rows)} VAT invoices | total invoice VAT={round(sum(r['vat'] for r in rows),2)} | suggested deductible={round(sum(r['ded'] for r in rows),2)}")

if __name__ == "__main__":
    main()
