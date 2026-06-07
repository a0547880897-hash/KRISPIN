"""Combine task-1 (pending drafts) + task-2 (approved expenses) into unified Excels.

Reads the two catalog xlsx detail sheets, normalises to one schema with a
'מקור' (source) column, and writes:
  combined_all_expenses.xlsx  - every expense document from both datasets
  combined_VAT_only.xlsx      - only invoices that carry deductible Israeli VAT
"""
import re
from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

GREEN = PatternFill("solid", fgColor="C6EFCE"); YELLOW = PatternFill("solid", fgColor="FFEB9C")
RED = PatternFill("solid", fgColor="FFC7CE"); HDR = PatternFill("solid", fgColor="305496")
SUB = PatternFill("solid", fgColor="DDEBF7"); THIN = Border(*[Side(style="thin", color="D9D9D9")]*4)

OUTCOLS = ["מקור", "#", "שם ספק", "ע.מ/ח.פ", "סוג מסמך", "מס' מסמך", "תאריך", "תקופה דו-חודשית",
           "סכום כולל מע\"מ", "מע\"מ", "מטבע", "סוג הוצאה", "סיווג PCN874", "ודאות", "הערה", "מזהה", "קובץ"]


def fnum(v):
    try: return round(float(v), 2)
    except Exception: return None


def bimonthly(d):
    m = re.search(r"(\d{4})-(\d{2})", str(d or ""))
    if not m: return ""
    y, mo = m.group(1), int(m.group(2)); s = mo if mo % 2 == 1 else mo-1
    return f"{s:02d}-{s+1:02d}/{y}"


def conf_norm(v):
    v = str(v or "")
    if "גבוה" in v or "🟩" in v: return "high"
    if "בינ" in v or "🟨" in v: return "medium"
    if "נמוכ" in v or "🟥" in v or "לבדיקה" in v: return "low"
    return ""


def load_task1():
    ws = load_workbook("catalog.xlsx")["קטלוג 441"]; h = [c.value for c in ws[1]]
    def col(n): return h.index(n)
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[col("#index")] in (None, ""): continue
        out.append({"מקור": "טיוטה ממתינה לאישור", "#": r[col("#index")], "שם ספק": r[col("שם ספק")],
                    "ע.מ/ח.פ": r[col("מספר עוסק / ח.פ.")], "סוג מסמך": r[col("סוג מסמך")],
                    "מס' מסמך": r[col("מספר מסמך")], "תאריך": r[col("תאריך המסמך")],
                    "תקופה דו-חודשית": r[col("תקופת דיווח")] or bimonthly(r[col("תאריך המסמך")]),
                    "סכום כולל מע\"מ": fnum(r[col("סכום כולל מע\"מ")]), "מע\"מ": fnum(r[col("סך הכל מע\"מ")]),
                    "מטבע": r[col("מטבע")], "סוג הוצאה": r[col("סוג הוצאה (מוצע)")],
                    "סיווג PCN874": r[col("סיווג PCN874")], "ודאות": conf_norm(r[col("רמת ודאות כללית")]),
                    "הערה": r[col("הערות")], "מזהה": r[col("draft_id")], "קובץ": r[col("שם הקובץ")]})
    return out


def load_task2():
    ws = load_workbook("catalog_expenses_FINAL.xlsx")["פירוט לפי חודש"]; h = [c.value for c in ws[1]]
    def col(n): return h.index(n)
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        v0 = r[col("#")]
        if v0 in (None, "") or str(r[col("expense_id")] or "") == "": continue  # skip month sub-headers
        out.append({"מקור": "מאושר/דווח", "#": v0, "שם ספק": r[col("שם ספק")],
                    "ע.מ/ח.פ": r[col("ע.מ/ח.פ")], "סוג מסמך": r[col("סוג מסמך")],
                    "מס' מסמך": r[col("מס' מסמך")], "תאריך": r[col("תאריך")],
                    "תקופה דו-חודשית": bimonthly(r[col("תאריך")]),
                    "סכום כולל מע\"מ": fnum(r[col("כולל מע\"מ")]), "מע\"מ": fnum(r[col("מע\"מ")]),
                    "מטבע": r[col("מטבע")], "סוג הוצאה": r[col("סוג הוצאה")],
                    "סיווג PCN874": r[col("סיווג PCN874")], "ודאות": conf_norm(r[col("ודאות")]),
                    "הערה": r[col("שאלה/ספק לליבון")], "מזהה": r[col("expense_id")], "קובץ": r[col("קובץ")]})
    return out


def write(rows, path, title):
    wb = Workbook(); ws = wb.active; ws.title = title; ws.sheet_view.rightToLeft = True
    for c, n in enumerate(OUTCOLS, 1):
        cell = ws.cell(1, c, n); cell.fill = HDR; cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = THIN
    ws.freeze_panes = "A2"
    cf = {"high": GREEN, "medium": YELLOW, "low": RED}
    for i, d in enumerate(rows, 2):
        for c, n in enumerate(OUTCOLS, 1):
            val = d.get(n, "")
            if n == "ודאות": val = {"high": "🟩 גבוהה", "medium": "🟨 בינונית", "low": "🟥 לבדיקה"}.get(d.get("ודאות"), "")
            cell = ws.cell(i, c, val); cell.border = THIN
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if n == "ודאות" and d.get("ודאות") in cf: cell.fill = cf[d["ודאות"]]
            if n == "מקור": cell.fill = SUB
    widths = [20, 6, 24, 13, 18, 14, 11, 12, 12, 10, 7, 20, 20, 12, 34, 22, 28]
    for c, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(c)].width = w
    wb.save(path)


t1, t2 = load_task1(), load_task2()
allrows = t1 + t2
write(allrows, "combined_all_expenses.xlsx", "כל המסמכים")
vat = [d for d in allrows if (d.get("מע\"מ") or 0) > 0 and (d.get("מטבע") or "ILS") == "ILS"]
write(vat, "combined_VAT_only.xlsx", "חשבוניות עם מע\"מ")
print(f"task1(טיוטות)={len(t1)} | task2(מאושר)={len(t2)} | combined={len(allrows)} | VAT-only={len(vat)}")
print("combined VAT total:", round(sum(d['מע\"מ'] for d in vat), 2))
